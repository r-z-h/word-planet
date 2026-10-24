"""把简历内容排成 A4 Word 文档（用于导出 PDF）。"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"D:\CodeProjects\Word Planet\简历-任自豪.docx"

CN_FONT = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x5B, 0x66, 0x72)


def set_font(run, size=10.2, bold=False, color=None, font=CN_FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def add_bottom_border(paragraph, size=6, color="1F3A5F"):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pPr.append(borders)


def rich(paragraph, text, size=10.2, color=None):
    """支持 **加粗** 的行内标记。"""
    for i, chunk in enumerate(text.split("**")):
        if not chunk:
            continue
        run = paragraph.add_run(chunk)
        set_font(run, size=size, bold=(i % 2 == 1), color=color)


def para(doc, text="", size=10.2, color=None, space_before=0, space_after=2,
         align=None, indent=None, hanging=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.06
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    if text:
        rich(p, text, size=size, color=color)
    return p


def bullet(doc, text, size=10.2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.55)
    pf.first_line_indent = Cm(-0.55)
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.06
    dot = p.add_run("• ")
    set_font(dot, size=size)
    rich(p, text, size=size)
    return p


def section(doc, title):
    p = para(doc, space_before=8, space_after=4)
    run = p.add_run(title)
    set_font(run, size=12, bold=True, color=ACCENT)
    add_bottom_border(p)


def main() -> None:
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.4)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.2)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    # ---------- 头部 ----------
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    run = p.add_run("任自豪")
    set_font(run, size=22, bold=True, color=ACCENT)

    para(doc, "求职意向：全栈开发 / Python 后端 / 前端开发",
         size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(doc, "13986804236　|　rzh0422@outlook.com　|　github.com/r-z-h",
         size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    para(doc, "武汉华夏理工学院 · 计算机科学与技术 · 本科大三",
         size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

    # ---------- 专业技能 ----------
    section(doc, "专业技能")
    bullet(doc, "**后端**：Python、FastAPI、Pydantic、SQLAlchemy 2.x、Alembic、PostgreSQL、SQLite（WAL）、JWT、APScheduler、RESTful API 设计与数据建模")
    bullet(doc, "**前端**：React 18、TypeScript、Vue 3、Pinia、Ant Design、Element Plus、ECharts、Vite、Axios")
    bullet(doc, "**测试**：pytest（接口 / 契约 / 边界）、Playwright（端到端）、真实链路冒烟、打包态验证")
    bullet(doc, "**工程**：金额精度治理、幂等设计、时区一致性、内容审计、软删除与审计日志、灰度式回归")

    # ---------- 项目一 ----------
    section(doc, "项目经历")

    p = para(doc, space_before=2, space_after=1)
    run = p.add_run("Word Planet · 少儿英语游戏化学习平台（个人项目，进行中）")
    set_font(run, size=11, bold=True)
    run2 = p.add_run("　　2026.10 — 至今")
    set_font(run2, size=9.5, color=MUTED)

    para(doc, "Vue 3 + FastAPI + PostgreSQL 全栈项目，正在把「人教版词表 → 分级 → 多套练习 → 错题复习 → 等级测试」串联成一条完整学习链路。", space_after=3)
    para(doc, "**规模**：21 个 REST 接口、11 张数据表、10 个页面、3105 道题（207 套）、后端约 1.7k 行 / 前端约 1.8k 行", space_after=4)

    para(doc, "**核心设计**", space_before=2, space_after=2)
    bullet(doc, "**三级内容模型**：知识点 → 题套 → 题目。7 个知识点共 207 套、每套固定 15 题；题套数据已生成完毕，按套推进与反复练习的调度在开发中。")
    bullet(doc, "**判题收口 + 幂等提交**：正确答案只存后端，前端仅上报作答内容；每次提交带幂等键，双击或网络重试不会重复计数、不会重复发放徽章。")
    bullet(doc, "**完成判定与防卡死**：由「必须全部答对」改为「每道题都做过一遍」，答错的题只进错题本、不回主流程，避免一道难题把整个知识点永久卡住。")
    bullet(doc, "**时间口径统一**：所有「今天」的统计显式按应用时区换算（AT TIME ZONE），并把同一口径用于每日任务、连续学习、家长报告与定时提醒，避免 UTC 服务器上的跨天错位。")
    bullet(doc, "**内容治理**：题目由 BSD-3 许可的人教版分年级词表（2618 条）清洗后生成，清洗管道丢弃约 3.5% 的脏数据并输出可抽查清单。")

    para(doc, "**质量保障**", space_before=3, space_after=2)
    bullet(doc, "三层验证：pytest 16 项接口测试 + 24 项真实 HTTP 冒烟 + 12 项专项回归；另写题库审计脚本，对 3105 道题校验选项完整性、答案与选项对应、听力题干扰词是否出现在原句、每套题量是否恰为 15，当前零问题。")
    bullet(doc, "修复一类「静默错题」：听力题的正确词必须出现在句子里、干扰项必须不出现，且不能只是句中某个词的一部分（如句中有 brothers、选项却是 brother）。")
    bullet(doc, "语音链路排障：Web Speech API 首次朗读有 3–4 秒延迟且会选到在线声音导致读音错误，改为「启动预热 + 本地声音打分优选」后，延迟与发音问题同时解决。")
    bullet(doc, "运维脚本：一条命令启停 PostgreSQL 与前后端，含端口占用识别（不误杀其他项目进程）、强杀残留 postmaster.pid 清理、就绪等待与日志落盘。")

    para(doc, "**当前进度**：核心学习闭环（分级测试 → 答题 → 错题本 → 解析 → 结算）已跑通并通过全部测试；按套推进、能力测试升级、多孩子档案、数据导出在开发中。", space_before=3, space_after=2)

    # ---------- 项目二 ----------
    p = para(doc, space_before=8, space_after=1)
    run = p.add_run("工程成本管理系统（独立设计与交付，真实用户在用）")
    set_font(run, size=11, bold=True)
    run2 = p.add_run("　　2026.10")
    set_font(run2, size=9.5, color=MUTED)

    para(doc, "FastAPI · SQLAlchemy · SQLite · React 18 · TypeScript · Ant Design · Playwright　GitHub：r-z-h/PIMS",
         size=9.5, color=MUTED, space_after=2)
    para(doc, "面向工程队老板的单机成本管理系统，覆盖成本录入、收付款、往来冲销、库存、结算、报表导出六大模块。后端 9.3k 行 / 前端 15.5k 行，82 个 REST 接口、16 张数据表、12 个页面，配 157 项接口测试 + 87 项端到端测试。",
         space_after=4)

    para(doc, "**核心设计**", space_before=2, space_after=2)
    bullet(doc, "**金额精度三层统一**：数据库 NUMERIC(14,2) → 后端 Decimal（统一 ROUND_HALF_UP）→ 前端整数「分」运算，全链路零浮点，避免对账差一分钱。")
    bullet(doc, "**消除「两套账」**：付款可从 4 个入口进入，原设计在结算单里另存「已发」导致对不上；改为已付一律由冲销记录现场汇总，四个入口自动一致，并把发放上限挂到同一口径上。")
    bullet(doc, "**自动冲销三级优先**：同对象 + 同类型 → 同类型 → 全量兜底，解决「给甲的钱冲了乙的欠款」。")
    bullet(doc, "**完工锁定**：项目完工后所有写操作需登录密码，解锁状态持久化到库（重启仍有效），只读接口不锁；被拦时前端自动弹窗、解锁后自动重试原操作。")
    bullet(doc, "**幂等与可迁移**：录成本自动生成应付单做了去重保护；旧库升级走「先迁移、后建表」的幂等脚本，避免 no such column 启动崩溃。")

    para(doc, "**质量保障**", space_before=3, space_after=2)
    bullet(doc, "建立契约测试：GET /projects/{id} 同时返回嵌套与扁平两种结构并锁死，修复了「mock 与真实实现不一致 → 测试全绿、页面整片空白」的假绿问题。")
    bullet(doc, "关键修复均做「能红一次」验证：故意改坏实现，确认测试确实失败，再还原。")
    bullet(doc, "复现并修复 20+ 个真实缺陷，含表头吸顶失效（CSS sticky 被 overflow 容器击败）、筛选入口重复、多入口状态不同步等。")
    bullet(doc, "文档：需求冻结、开发手册、数据库设计、接口设计、体系结构图、自测交接说明共 6 份。")

    # ---------- 教育 ----------
    section(doc, "教育经历")
    bullet(doc, "**武汉华夏理工学院**　计算机科学与技术　本科　大三")

    # ---------- 其他 ----------
    section(doc, "证书与其他")
    bullet(doc, "普通话二级乙等　·　C1 驾驶证　·　计算机一级")
    bullet(doc, "有餐饮行业暑期工经历")

    section(doc, "自我评价")
    para(doc, "从 0 到 1 主导项目的落地与迭代，主攻 AI Agent 方向，具备全栈开发能力。学习能力强，能吃苦，有团队意识，沟通与表达能力较好。")

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
